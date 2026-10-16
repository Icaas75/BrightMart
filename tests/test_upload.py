"""Tests for the upload-driven flow: file in -> report generated."""

from __future__ import annotations

import datetime as dt
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest

from src.dashboard import (
    DEFAULT_DATA,
    UPLOAD_DIR,
    app,
    clear_report_cache,
)

HEADER = [
    "Date", "Invoice_ID", "Salesperson", "Product", "Category",
    "Quantity", "Unit_Price", "Discount", "Payment_Method", "Customer_Type",
]


def _row(date: dt.datetime, invoice: str, qty: int, price: float, disc: float) -> list:
    return [
        date, invoice, "Sara", "USB-C Cable", "Accessories",
        qty, price, disc, "Cash", "Individual",
    ]


@pytest.fixture()
def fresh_app():
    """No data loaded; upload dir cleaned after the test."""
    app.config["DATA_PATH"] = None
    app.config["SOURCE_NAME"] = ""
    app.config["SHEET"] = None
    clear_report_cache()
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    yield app
    for old in UPLOAD_DIR.glob("upload-*.xlsx"):
        old.unlink(missing_ok=True)
    app.config["DATA_PATH"] = None
    app.config["SOURCE_NAME"] = ""
    clear_report_cache()


def _client(fresh_app):
    return fresh_app.test_client()


def _upload(client, content: bytes, filename: str = "month.xlsx"):
    # The payload MUST be a file-like object: with raw bytes the Werkzeug 3.x
    # test client sends it as a form field and request.files parses empty.
    return client.post(
        "/upload",
        data={"file": (BytesIO(content), filename)},
    )


def _workbook_bytes(rows: list[list]) -> bytes:
    """Build a workbook in memory; rows[0] is the header."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Data"
    for row in rows:
        ws.append(row)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _valid_rows() -> list[list]:
    return [
        HEADER,
        _row(dt.datetime(2026, 9, 3, 10, 0), "INV-A", 2, 900, 0),    # 1800
        _row(dt.datetime(2026, 9, 5, 11, 0), "INV-B", 1, 1800, 0.1), # 1620
    ]


# ------------------------------------------------------------- happy path


def test_upload_generates_report(fresh_app):
    client = _client(fresh_app)
    response = _upload(client, _workbook_bytes(_valid_rows()))
    assert response.status_code == 302

    page = client.get("/").get_data(as_text=True)
    assert "3,420.00" in page              # 2*900*1 + 1*1800*0.9
    assert "Sep 03, 2026" in page
    assert "Source: month.xlsx" in page


def test_upload_then_export_matches(fresh_app):
    client = _client(fresh_app)
    _upload(client, _workbook_bytes(_valid_rows()))
    response = client.get("/export.xlsx")
    assert response.status_code == 200
    assert response.data[:2] == b"PK"
    wb = openpyxl.load_workbook(__import__("io").BytesIO(response.data))
    assert "Summary" in wb.sheetnames


def test_second_upload_replaces_first(fresh_app):
    client = _client(fresh_app)
    _upload(client, _workbook_bytes(_valid_rows()))
    october = [
        HEADER,
        _row(dt.datetime(2026, 10, 7, 9, 0), "INV-C", 1, 5000, 0),  # 5000
    ]
    _upload(client, _workbook_bytes(october), filename="october.xlsx")
    page = client.get("/").get_data(as_text=True)
    assert "5,000.00" in page
    assert "Oct 07, 2026" in page
    assert "Source: october.xlsx" in page


def test_use_default_sample(fresh_app):
    if not DEFAULT_DATA.is_file():
        pytest.skip("sample file not present in project root")
    client = _client(fresh_app)
    response = client.post("/use-default")
    assert response.status_code == 302
    page = client.get("/").get_data(as_text=True)
    assert "26,351,224.00" in page          # the August sample totals
    assert "Source: brightmart_sales_data.xlsx" in page


# ------------------------------------------------------------- rejections


def test_upload_non_excel_rejected(fresh_app):
    client = _client(fresh_app)
    response = _upload(client, b"not an excel file", filename="notes.csv")
    assert response.status_code == 400
    assert "not an Excel workbook" in response.get_data(as_text=True)
    # dashboard still has no data
    assert client.get("/").status_code == 200  # upload page again
    assert b"svg" not in client.get("/").data


def test_upload_corrupt_xlsx_rejected_with_reason(fresh_app):
    client = _client(fresh_app)
    response = _upload(client, b"PK fake zip content", filename="broken.xlsx")
    assert response.status_code == 400
    assert "Could not use that file" in response.get_data(as_text=True)


def test_upload_missing_column_rejected(fresh_app):
    client = _client(fresh_app)
    bad_header = [c for c in HEADER if c != "Unit_Price"]
    rows = [bad_header, _row(dt.datetime(2026, 9, 1), "INV-1", 1, 900, 0)[:-1]]
    response = _upload(client, _workbook_bytes(rows))
    assert response.status_code == 400
    assert "Missing required columns" in response.get_data(as_text=True)
    # Jinja escapes the quotes around the column name, so match it bare
    assert "Unit_Price" in response.get_data(as_text=True)


def test_upload_empty_sheet_rejected(fresh_app):
    client = _client(fresh_app)
    response = _upload(client, _workbook_bytes([HEADER]))
    assert response.status_code == 400
    assert "no data rows" in response.get_data(as_text=True).lower()


def test_upload_no_file_selected(fresh_app):
    client = _client(fresh_app)
    response = client.post("/upload", data={})
    assert response.status_code == 400
    assert "No file selected" in response.get_data(as_text=True)


def test_upload_page_shown_when_no_data(fresh_app):
    client = _client(fresh_app)
    page = client.get("/")
    assert page.status_code == 200
    assert "Drop an .xlsx here" in page.get_data(as_text=True)


def test_export_before_any_upload_redirects_to_upload(fresh_app):
    client = _client(fresh_app)
    response = client.get("/export.xlsx")
    assert response.status_code == 302
