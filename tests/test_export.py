"""Tests for src.export: the Excel output stage."""

from __future__ import annotations

import datetime as dt
import math
from io import BytesIO
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest
from openpyxl.chart import BarChart, DoughnutChart, LineChart

from src.export import (
    _clean_value,
    export_report,
    main,
    report_to_bytes,
    write_report_workbook,
)
from src.loader import REQUIRED_COLUMNS, load_sales
from src.report import build_report

EXPECTED_SHEETS = [
    "Summary",
    "Sales Dashboard",
    "Daily Revenue",
    "Top Products",
    "By Category",
    "By Salesperson",
    "By Payment Method",
    "By Customer Type",
    "Cleaned Data",
    "Flagged for Review",
    "Chart Data",
]


def _row(**over) -> dict:
    values = {
        "Date": dt.datetime(2026, 8, 1, 9, 0),
        "Invoice_ID": "INV-1",
        "Salesperson": "Sara",
        "Product": "USB-C Cable",
        "Category": "Accessories",
        "Quantity": 2,
        "Unit_Price": 900,
        "Discount": 0,
        "Payment_Method": "Cash",
        "Customer_Type": "Individual",
    }
    values.update(over)
    return values


# ------------------------------------------------------------- real workbook


def test_real_report_workbook_sheets(real_sales_file):
    wb = write_report_workbook(build_report(real_sales_file))
    assert wb.sheetnames == EXPECTED_SHEETS


def test_real_summary_sheet_kpis(real_sales_file):
    report = build_report(real_sales_file)
    wb = write_report_workbook(report)
    ws = wb["Summary"]

    labels = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
              for r in range(1, ws.max_row + 1)}
    assert labels["Total revenue (ETB)"] == pytest.approx(
        report.totals["total_revenue"]
    )
    assert labels["Orders"] == report.totals["orders"]
    assert labels["Units sold"] == report.totals["units"]
    assert labels["Rows rows in"] == report.cleaning_summary["rows_in"]
    assert labels["Rows flagged for review"] == report.cleaning_summary["flagged_for_review"]


def test_real_daily_sheet_matches_report(real_sales_file):
    report = build_report(real_sales_file)
    wb = write_report_workbook(report)
    ws = wb["Daily Revenue"]

    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    # enriched columns are appended after the raw aggregation columns
    assert header[: len(report.daily.columns)] == list(report.daily.columns)
    assert header[len(report.daily.columns):] == ["7-Day Avg", "Cumulative"]
    assert ws.max_row - 1 == len(report.daily) == 31

    # first data row: the earliest day
    assert ws.cell(row=2, column=1).value == report.daily.loc[0, "Date"].to_pydatetime()
    assert ws.cell(row=2, column=2).value == pytest.approx(
        report.daily.loc[0, "Revenue"]
    )


def test_real_cleaned_and_flagged_sheets(real_sales_file):
    report = build_report(real_sales_file)
    wb = write_report_workbook(report)

    cleaned = wb["Cleaned Data"]
    assert cleaned.max_row - 1 == len(report.rows) == 498
    # raw columns + the computed Revenue column
    assert cleaned.max_column == len(report.rows.columns) + 1

    flagged = wb["Flagged for Review"]
    assert flagged.max_row - 1 == len(report.flagged) == 2


def test_breakdown_sheets_match_report(real_sales_file):
    report = build_report(real_sales_file)
    wb = write_report_workbook(report)

    by_cat = wb["By Category"]
    header = [by_cat.cell(row=1, column=c).value for c in range(1, by_cat.max_column + 1)]
    # enriched layout: label, revenue, share, orders
    assert header == ["Category", "Total_Revenue", "Share_of_Total", "Orders"]
    assert by_cat.max_row - 1 == len(report.by_category) == 4
    assert by_cat.cell(row=2, column=1).value == report.by_category.loc[0, "Category"]
    assert by_cat.cell(row=2, column=2).value == pytest.approx(
        report.by_category.loc[0, "Total_Revenue"]
    )


# ------------------------------------------------------- bytes serialization


def test_report_to_bytes_produces_valid_workbook(real_sales_file):
    report = build_report(real_sales_file)
    data = report_to_bytes(report)
    assert isinstance(data, bytes) and len(data) > 0

    wb = openpyxl.load_workbook(BytesIO(data))
    assert wb.sheetnames == EXPECTED_SHEETS


def test_bytes_export_is_deterministic_content(real_sales_file):
    """Same report -> same data (ignoring the generated-at timestamp)."""
    report = build_report(real_sales_file)
    wb = write_report_workbook(report)
    again = write_report_workbook(report)
    assert (
        wb["Top Products"]["B2"].value == again["Top Products"]["B2"].value
    )


# -------------------------------------------------------------- file writing


def test_export_report_writes_file(real_sales_file, tmp_path):
    report = build_report(real_sales_file)
    target = tmp_path / "report.xlsx"
    result = export_report(report, target)

    assert result == target
    assert target.is_file()
    wb = openpyxl.load_workbook(target)
    assert wb.sheetnames == EXPECTED_SHEETS


# ---------------------------------------------------------------------- CLI


def test_cli_writes_default_output(real_sales_file, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    assert main(["export", str(real_sales_file)]) == 0
    assert (tmp_path / "brightmart_sales_report.xlsx").is_file()


def test_cli_custom_output(real_sales_file, tmp_path):
    out = tmp_path / "custom.xlsx"
    assert main(["export", str(real_sales_file), str(out)]) == 0
    assert out.is_file()


def test_cli_bad_usage_returns_2():
    assert main(["export"]) == 2
    assert main(["export", "a", "b", "c"]) == 2


# ------------------------------------------------------------- value helper


def test_clean_value_handles_scalars():
    assert _clean_value(np.int64(7)) == 7
    assert _clean_value(np.float64(2.5)) == 2.5
    stamp = pd.Timestamp("2026-08-01 09:00")
    assert _clean_value(stamp) == dt.datetime(2026, 8, 1, 9, 0)
    assert _clean_value(float("nan")) is None
    assert _clean_value(pd.NaT) is None
    assert _clean_value(None) is None
    assert _clean_value("text") == "text"
